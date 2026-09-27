from __future__ import annotations

import csv
import textwrap
import uuid
from io import BytesIO, StringIO
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session
from sqlalchemy import or_

from app.core.config import settings
from app.core.auth import require_roles
from app.core.database import get_db
from app.models import Document, DocCategory, DocSource, DocStatus, FieldExtraction, Role

router = APIRouter(tags=["documents"])

CATEGORY_MAP = {
    "PROSPECTING": DocCategory.PROSPECTING,
    "GEOLOGICAL_REPORT": DocCategory.GEOLOGICAL_REPORT,
    "PRODUCTION": DocCategory.PRODUCTION,
    "ENVIRONMENTAL": DocCategory.ENVIRONMENTAL,
    "RESERVES": DocCategory.RESERVES,
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
SHEET_EXTENSIONS = {".xlsx", ".csv"}
MAX_UPLOAD_BYTES = 50 * 1024 * 1024


def _spreadsheet_rows(filename: str, content: bytes) -> list[tuple[str, list[list[str]]]]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
            rows = [[str(cell) for cell in row] for row in csv.reader(StringIO(text))]
        except (UnicodeDecodeError, csv.Error) as exc:
            raise HTTPException(415, f"{filename}: CSV must be valid UTF-8") from exc
        return [(Path(filename).stem, rows)]

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in workbook.worksheets:
                rows = [
                    ["" if value is None else str(value) for value in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                sheets.append((sheet.title, rows))
            return sheets
        finally:
            workbook.close()
    except Exception as exc:
        raise HTTPException(415, f"{filename}: could not read XLSX workbook") from exc


def _render_spreadsheet_pdf(filename: str, content: bytes, destination: Path) -> None:
    import fitz

    sheets = _spreadsheet_rows(filename, content)
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    y = 48

    def write_line(line: str, *, bold: bool = False) -> None:
        nonlocal page, y
        if y > 790:
            page = pdf.new_page(width=595, height=842)
            y = 48
        page.insert_text((48, y), line[:160], fontsize=10, fontname="hebo" if bold else "helv")
        y += 14

    for sheet_name, rows in sheets:
        write_line(f"Worksheet: {sheet_name}", bold=True)
        if len(rows) > 1 and any(cell.strip() for cell in rows[0]):
            headers = [cell.strip() or f"Column {index + 1}" for index, cell in enumerate(rows[0])]
            for row in rows[1:]:
                for index, value in enumerate(row):
                    if index >= len(headers) or not value.strip():
                        continue
                    labeled = f"{headers[index]}: {value.strip()}"
                    for line in textwrap.wrap(labeled, width=105, break_long_words=True) or [""]:
                        write_line(line)
                y += 6
        else:
            for row in rows:
                flattened = " | ".join(cell.replace("\n", " ").strip() for cell in row).strip()
                if not flattened:
                    continue
                for line in textwrap.wrap(flattened, width=105, break_long_words=True) or [""]:
                    write_line(line)
        y += 12
    if not pdf.page_count:
        pdf.new_page()
    pdf.save(destination)
    pdf.close()


def _normalize_upload(filename: str, content: bytes, destination: Path) -> DocSource:
    """Validate uploads and convert supported images/sheets to a PDF for ingestion."""
    extension = Path(filename).suffix.lower()
    if not content:
        raise HTTPException(400, f"{filename}: file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"{filename}: maximum upload size is 50 MB")
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise HTTPException(415, f"{filename}: file does not contain a valid PDF")
        try:
            import fitz
            with fitz.open(stream=content, filetype="pdf") as source:
                if source.page_count == 0:
                    raise ValueError("PDF has no pages")
        except Exception as exc:
            raise HTTPException(415, f"{filename}: could not read PDF") from exc
        destination.write_bytes(content)
        return DocSource.PDF

    if extension in IMAGE_EXTENSIONS:
        try:
            from PIL import Image
            with Image.open(BytesIO(content)) as image:
                image.seek(0)
                image.convert("RGB").save(destination, format="PDF", resolution=150)
        except Exception as exc:
            raise HTTPException(415, f"{filename}: unsupported or corrupt image") from exc
        return DocSource.IMAGE

    if extension in SHEET_EXTENSIONS:
        _render_spreadsheet_pdf(filename, content, destination)
        return DocSource.SPREADSHEET

    raise HTTPException(415, f"{filename}: supported formats are PDF, images, XLSX, and CSV")


def _doc_payload(d: Document, fields_count: int = 0) -> dict:
    c = d.coalfield
    return {
        "id": d.id,
        "title": d.title,
        "category": d.category.value if d.category else "GEOLOGICAL_REPORT",
        "coalfield_id": d.coalfield_id or "",
        "coalfield_name": c.name if c else "",
        "block": d.block_name or "",
        "district": d.district or "",
        "source_type": d.source_type.value if d.source_type else "PDF",
        "status": d.status.value if d.status else "UPLOADING",
        "ingestion_progress": d.ingestion_progress or {"step": "UPLOADING", "pct": 0, "message": ""},
        "page_count": d.page_count or 0,
        "size_bytes": d.size_bytes or 0,
        "total_confidence": d.total_confidence or 0,
        "report_date": d.report_date or "",
        "uploaded_at": d.uploaded_at.isoformat() if d.uploaded_at else "",
        "pages_dir": d.pages_dir or "",
    }


@router.get("/documents")
def list_documents(category: str | None = None, coalfield: str | None = None,
                   status: str | None = None, q: str | None = None, db: Session = Depends(get_db)):
    query = db.query(Document)
    if category:
        query = query.filter(Document.category == CATEGORY_MAP.get(category))
    if coalfield:
        from app.models import Coalfield
        query = query.join(Coalfield, Document.coalfield_id == Coalfield.id).filter(Coalfield.name == coalfield)
    if status:
        query = query.filter(Document.status == status)
    if q:
        query = query.filter(or_(
            Document.title.ilike(f"%{q}%"),
            Document.block_name.ilike(f"%{q}%"),
        ))
    docs = query.order_by(Document.uploaded_at.desc()).all()
    return [_doc_payload(d) for d in docs]


@router.get("/documents/coalfields")
def list_coalfields(db: Session = Depends(get_db)):
    from app.models import Coalfield
    return [
        {"id": c.id, "name": c.name, "subsidiary": c.subsidiary, "state": c.state,
         "area_sqkm": getattr(c, "area_sqkm", 0) or 0,
         "baseline_gcv": getattr(c, "baseline_gcv", 0) or 0,
         "baseline_ash": getattr(c, "baseline_ash", 0) or 0,
         "baseline_ob_ratio": getattr(c, "baseline_ob_ratio", 0) or 0,
         "reserve_mt": getattr(c, "reserve_mt", 0) or 0}
        for c in db.query(Coalfield).all()
    ]


@router.get("/documents/{doc_id}")
def get_document(doc_id: str, db: Session = Depends(get_db)):
    d = db.query(Document).get(doc_id)
    if d is None:
        raise HTTPException(404, "Document not found")
    n = db.query(FieldExtraction).filter(FieldExtraction.document_id == d.id).count()
    return _doc_payload(d, n)


@router.get("/documents/{doc_id}/status")
def doc_status(doc_id: str, db: Session = Depends(get_db)):
    d = db.query(Document).get(doc_id)
    if d is None:
        raise HTTPException(404, "Document not found")
    return {
        "document_id": d.id,
        "status": d.status.value if d.status else "UPLOADING",
        "step": (d.ingestion_progress or {}).get("step", ""),
        "pct": (d.ingestion_progress or {}).get("pct", 0),
        "message": (d.ingestion_progress or {}).get("message", ""),
    }


@router.get("/documents/{doc_id}/progress")
def doc_progress(doc_id: str, db: Session = Depends(get_db)):
    """Legacy alias."""
    return doc_status(doc_id, db)


def _run_ingest(doc_id: str):
    from app.services.ingestion import ingest_document
    ingest_document(doc_id)


@router.post("/documents/upload", dependencies=[Depends(require_roles(Role.GEOLOGIST))])
async def upload_documents(
    background: BackgroundTasks,
    files: list[UploadFile] = File(...),
    title: str | None = None,
    category: str = "GEOLOGICAL_REPORT",
    as_user: str = "GEOLOGIST",
    db: Session = Depends(get_db),
):
    upload_root = Path(settings.uploads_dir)
    upload_root.mkdir(parents=True, exist_ok=True)

    created = []
    for f in files:
        name = Path(f.filename or "doc.pdf").name
        stem = Path(name).stem[:100] or "document"
        dest = upload_root / f"{stem}-{uuid.uuid4().hex[:8]}.pdf"
        content = await f.read()
        source_type = _normalize_upload(name, content, dest)
        doc = Document(
            title=title or name,
            category=CATEGORY_MAP.get(category, DocCategory.GEOLOGICAL_REPORT),
            source_type=source_type,
            status=DocStatus.UPLOADING,
            size_bytes=len(content),
            file_path=str(dest),
            ingestion_progress={"step": "UPLOADING", "pct": 5, "message": "Queued"},
        )
        db.add(doc)
        db.flush()
        created.append(doc)
        db.commit()
        background.add_task(_run_ingest, doc.id)

    # Frontend expects bare Document[]
    return [_doc_payload(d) for d in created]


@router.get("/documents/{doc_id}/file")
def doc_file(doc_id: str, db: Session = Depends(get_db)):
    d = db.query(Document).get(doc_id)
    if d is None or not d.file_path:
        raise HTTPException(404, "Source file unavailable")
    file_path = Path(d.file_path)
    if not file_path.is_absolute():
        file_path = Path(__file__).resolve().parents[2] / file_path
    if not file_path.is_file():
        raise HTTPException(404, "Source file unavailable")
    from fastapi.responses import FileResponse
    return FileResponse(file_path, filename=f"{d.title or 'doc'}.pdf")


@router.delete("/documents/{doc_id}", dependencies=[Depends(require_roles(Role.GEOLOGIST))])
def delete_document(doc_id: str, db: Session = Depends(get_db)):
    d = db.query(Document).get(doc_id)
    if d is None:
        raise HTTPException(404, "Document not found")
    db.delete(d)
    db.commit()
    return {"ok": True, "id": doc_id}
